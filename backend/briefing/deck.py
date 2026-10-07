"""
briefing/deck.py — the issue as a deck, for the meeting where it is read.

Built from the same document as the PDF, so the two never disagree: a cover,
one slide per key judgment, our sites, the exposure register, the decisions,
one slide per section (its lead and what it means for the reader), the
scenarios, our imagery and what to watch. 16:9, the report's palette.
"""
from __future__ import annotations

import base64
import io
import re
from pathlib import Path

from . import render, spec

INK = (0x1B, 0x1F, 0x24)
MUTED = (0x6B, 0x6B, 0x6B)
ACCENT = (0x2F, 0x5C, 0x90)
RULE = (0xC9, 0xC5, 0xBC)
LEVEL = {1: (0x5F, 0x7A, 0x5C), 2: (0xA8, 0x9F, 0x68), 3: (0xB8, 0x89, 0x2F), 4: (0xA8, 0x5F, 0x3A), 5: (0x8B, 0x3A, 0x2F)}


def _plain(t) -> str:
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", str(t or ""))
    return re.sub(r"\s+", " ", t).strip()


def build(doc: dict, out: str | Path) -> str:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt

    L = render.labels(doc["meta"].get("language", "de"))
    m = doc["meta"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W = prs.slide_width

    def rgb(c):
        return RGBColor(*c)

    def text(slide, x, y, w, h, s, size=14, bold=False, color=INK, font="Inter"):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        paras = s if isinstance(s, list) else [s]
        for i, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = _plain(ptxt)
            p.font.size, p.font.bold, p.font.name = Pt(size), bold, font
            p.font.color.rgb = rgb(color)
            p.space_after = Pt(6)
        return tb

    def frame(title, kicker=""):
        s = prs.slides.add_slide(blank)
        if kicker:
            text(s, 0.6, 0.35, 10, 0.3, kicker.upper(), 10, True, ACCENT)
        text(s, 0.6, 0.6, 12, 0.8, title, 26, True)
        line = s.shapes.add_connector(1, Inches(0.6), Inches(1.4), W - Inches(0.6), Inches(1.4))
        line.line.color.rgb = rgb(INK)
        text(s, 0.6, 7.0, 8, 0.3, f"PARALLAX · {m.get('serial', '')}", 8, False, MUTED)
        if m.get("rehearsal"):
            text(s, 9.0, 7.0, 3.8, 0.3, L["rehearsal"], 8, True, (0x8B, 0x3A, 0x2F))
        return s

    def table(slide, cols, rows, y=1.7, widths=None, size=11):
        n = len(rows) + 1
        h = min(5.0, 0.42 * n)
        shp = slide.shapes.add_table(n, len(cols), Inches(0.6), Inches(y), W - Inches(1.2), Inches(h))
        t = shp.table
        if widths:
            total = sum(widths)
            for i, w in enumerate(widths):
                t.columns[i].width = Emu(int((W - Inches(1.2)) * w / total))
        for j, c in enumerate(cols):
            cell = t.cell(0, j)
            cell.text = c
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(INK)
            for p in cell.text_frame.paragraphs:
                p.font.size, p.font.bold, p.font.name = Pt(size), True, "Inter"
                p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for i, r in enumerate(rows, 1):
            for j, v in enumerate(r):
                cell = t.cell(i, j)
                cell.text = _plain(v)
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i % 2 else RGBColor(0xF4, 0xF1, 0xEB)
                for p in cell.text_frame.paragraphs:
                    p.font.size, p.font.name = Pt(size - 1), "Inter"
                    p.font.color.rgb = rgb(INK)
        return t

    # cover
    s = prs.slides.add_slide(blank)
    text(s, 0.8, 0.8, 8, 0.5, "PARALLAX", 20, True, INK)
    text(s, 0.8, 1.25, 8, 0.4, (m.get("publisher_line") or "").upper(), 10, False, MUTED)
    text(s, 0.8, 2.6, 11.5, 1.2, L[m["cadence"]], 40, True)
    text(s, 0.8, 3.7, 11.5, 0.6, f"{L['for']} {m.get('org', '')}", 18, False, INK, "Source Serif 4")
    text(s, 0.8, 4.3, 11.5, 0.5, f"{L['period']} {m.get('period_label', '')}", 14, False, MUTED)
    text(s, 0.8, 6.3, 11.5, 0.5, f"{L['serial']} {m.get('serial', '')} · {L['cutoff']} {m.get('cutoff_label', '')}", 10, False, MUTED)
    if m.get("rehearsal"):
        text(s, 0.8, 5.0, 11.5, 0.4, L["rehearsal"], 12, True, (0x8B, 0x3A, 0x2F))

    for k in doc.get("key_judgments") or []:
        s = frame(k.get("title", ""), f"{L['key_judgments']} · {k.get('id', '')}")
        text(s, 0.6, 1.7, 12, 3.6, k.get("body", ""), 16, False, INK, "Source Serif 4")
        band = spec.BAND_LABEL[m.get("language", "de")].get(k.get("band"), "")
        text(s, 0.6, 5.7, 4, 0.8, [L["probability"].upper(), f"{band} ({k.get('low')}–{k.get('high')} %)"], 12, True)
        text(s, 4.8, 5.7, 3, 0.8, [L["confidence"].upper(), L["conf"].get(k.get("confidence"), "")], 12, True)
        text(s, 8.2, 5.7, 4, 0.8, [L["change"].upper(), k.get("change", "")], 12, True)

    if doc.get("sites"):
        s = frame(L["sites"])
        table(s, [L["site"], L["kind"], L["at_site"], L["precedents"], L["strongest"]],
              [[x["name"], x.get("kind_label", ""), str(x.get("events", 0)), str(x.get("precedents", 0)),
                (x.get("strongest") or {}).get("title", L["none_short"])] for x in doc["sites"]], widths=[3, 2, 1.2, 1.2, 5])

    ex = doc.get("exposure")
    if ex and ex.get("vectors"):
        s = frame(L["exposure"])
        t = table(s, ["", L["vector"], L["level"], L["prev"], "Δ", L["drivers"]],
                  [[v.get("id", ""), v.get("name", ""), str(v.get("level", "")), str(v.get("prev", "")), v.get("delta", ""), v.get("drivers", "")]
                   for v in ex["vectors"]], widths=[0.6, 3, 0.8, 0.9, 0.6, 6.5], size=10)
        for i, v in enumerate(ex["vectors"], 1):
            c = LEVEL.get(int(v.get("level") or 1) if str(v.get("level", "")).isdigit() else 1)
            cell = t.cell(i, 2)
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(c)
            for p in cell.text_frame.paragraphs:
                p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); p.font.bold = True

    de = doc.get("decisions")
    if de and de.get("items"):
        s = frame(L["decisions"])
        table(s, ["", L["decide"], L["owner"], L["due"], L["urgency"]],
              [[d.get("id", ""), d.get("what", ""), d.get("owner", ""), d.get("due", ""), d.get("urgency", "")] for d in de["items"]],
              widths=[0.7, 6, 2.4, 1.8, 1.4])

    for sec in doc.get("sections") or []:
        s = frame(sec.get("title", ""), sec.get("kicker", ""))
        meaning = next((b["text"] for sub in reversed(sec.get("subsections") or [{"blocks": sec.get("blocks") or []}])
                        for b in reversed(sub["blocks"]) if b.get("type") == "meaning"), "")
        if not meaning:
            meaning = next((b["text"] for b in reversed(sec.get("blocks") or []) if b.get("type") == "meaning"), "")
        text(s, 0.6, 1.7, 12, 1.0, sec.get("lead", ""), 18, False, INK, "Source Serif 4")
        subs = [sub.get("title", "") for sub in sec.get("subsections") or []]
        if subs:
            text(s, 0.6, 2.9, 5.5, 3.5, subs, 13, False, MUTED)
        box = s.shapes.add_shape(1, Inches(6.4), Inches(2.9), Inches(6.3), Inches(3.6))
        box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xF4, 0xF1, 0xEB); box.line.color.rgb = rgb(RULE)
        text(s, 6.6, 3.0, 6.0, 0.4, f"{L['meaning_for'].upper()} {m.get('org', '').upper()}", 10, True, ACCENT)
        text(s, 6.6, 3.4, 6.0, 3.0, re.sub(r"\s*\[[^\]]+\]", "", meaning), 13, False, INK, "Source Serif 4")

    sc = doc.get("scenarios")
    if sc and sc.get("items"):
        s = frame(L["scenarios"])
        table(s, ["", "", "%", L["trigger"]], [[x.get("key", ""), x.get("title", ""), str(x.get("p", "")), x.get("trigger", "")]
                                              for x in sc["items"]], widths=[0.5, 4, 0.7, 7])

    im = doc.get("imagery")
    for it in (im or {}).get("items") or []:
        s = frame(L["imagery"])
        try:
            raw = base64.b64decode(it["src"].split(",", 1)[1])
            s.shapes.add_picture(io.BytesIO(raw), Inches(0.6), Inches(1.7), height=Inches(5.0))
        except Exception:                                    # noqa: BLE001
            pass
        text(s, 6.4, 1.7, 6.3, 4, re.sub(r"\s*\[[^\]]+\]", "", it.get("caption", "")), 13, False, INK, "Source Serif 4")

    ind = doc.get("indicators")
    if ind and ind.get("rows"):
        s = frame(L["indicators"] if m["cadence"] != "daily" else L["watch"])
        table(s, ["", L["indicators_col"], L["threshold"], L["current"], L["trend"]],
              [[r.get("id", ""), r.get("name", ""), r.get("threshold", ""), r.get("current", ""), r.get("trend", "")] for r in ind["rows"][:10]],
              widths=[0.7, 6, 2, 1.6, 1.4], size=10)

    out = Path(out)
    prs.save(str(out))
    return str(out)
